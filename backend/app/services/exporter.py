import csv
import io
import re

from docx import Document
from docx.shared import Pt, RGBColor

from app.schemas import AnalysisSession, RequirementResult

_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def final_text(r: RequirementResult) -> str:
    """The requirement as it should appear in the corrected document."""
    if r.decision in ("accepted", "edited") and r.final_text:
        return r.final_text
    return r.original


def _safe(v: str) -> str:
    """Neutralise spreadsheet formula injection (requirements are untrusted text)."""
    return "'" + v if v.startswith(_FORMULA) else v


def filename(s: AnalysisSession, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", s.name).strip("_.") or "analysis"
    return f"{base[:60]}.{ext}"


def to_json(s: AnalysisSession) -> bytes:
    return s.model_dump_json(indent=2).encode()


def to_csv(s: AnalysisSession) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["id", "original", "ambiguous", "issues", "max_severity", "suggested_rewrite", "decision",
                "final_text", "clarifying_questions", "quality_before", "quality_after", "confidence"])
    order = {"high": 3, "medium": 2, "low": 1}
    for r in s.requirements:
        top = max((i.severity for i in r.issues), key=order.get, default="")
        w.writerow([r.id, _safe(r.original), "yes" if r.is_ambiguous else "no",
                    _safe("; ".join(f"{i.type}: {i.span}" for i in r.issues)), top,
                    _safe(r.suggested_rewrite or ""), r.decision, _safe(final_text(r)),
                    _safe(" | ".join(r.clarifying_questions)), r.quality_before, r.quality_after, r.confidence])
    return ("﻿" + buf.getvalue()).encode()  # BOM so Excel reads UTF-8


def to_docx(s: AnalysisSession) -> bytes:
    d = Document()
    d.add_heading(s.name, 0)
    st = s.summary
    d.add_paragraph(
        f"{st.total} requirements analysed · {st.ambiguous} ambiguous ({st.ambiguity_rate:.0%}) · "
        f"average quality {st.avg_quality_before:.0f} → {st.avg_quality_after:.0f}")

    d.add_heading("Corrected requirements", 1)
    for r in s.requirements:
        p = d.add_paragraph(style="List Number")
        p.add_run(f"[{r.id}] ").bold = True
        run = p.add_run(final_text(r))
        if r.is_ambiguous and r.decision == "pending":
            run.font.color.rgb = RGBColor(0xB4, 0x5F, 0x06)
            p.add_run("  (unresolved)").italic = True

    d.add_heading("Analysis details", 1)
    open_items = [r for r in s.requirements if r.is_ambiguous]
    if not open_items:
        d.add_paragraph("No ambiguous requirements were found.")
    for r in open_items:
        d.add_heading(f"{r.id} — {r.decision}", 2)
        d.add_paragraph().add_run("Original: ").bold = True
        d.paragraphs[-1].add_run(r.original)
        t = d.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ("Type", "Severity", "Text", "Why")):
            c.text = h
        for i in r.issues:
            row = t.add_row().cells
            row[0].text, row[1].text, row[2].text, row[3].text = i.type.replace("_", " "), i.severity, i.span, i.explanation
        if r.suggested_rewrite:
            d.add_paragraph().add_run("Suggested rewrite: ").bold = True
            d.paragraphs[-1].add_run(r.suggested_rewrite)
        for q in r.clarifying_questions:
            d.add_paragraph(q, style="List Bullet")
    for p in d.paragraphs:
        for run in p.runs:
            run.font.size = run.font.size or Pt(11)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


EXPORTERS = {
    "json": (to_json, "application/json"),
    "csv": (to_csv, "text/csv; charset=utf-8"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
}
